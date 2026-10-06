"""Configuration handling and the JSON the front end consumes."""

import base64
import copy
import json

import numpy as np
import pytest

from server import report_service, service, twin_service, validation_service

LEGACY = {  # a scenario file written by the Streamlit versions of Taki
    "schema_version": 1, "num_lines": 2, "max_sag_m": 2.0, "row_boundary_m": 15.0,
    "lines": [
        {"preset_name": "132kV Double Circuit Lattice", "x_offset": -15, "load": 60, "arrangement": "ABC-CBA",
         "current_override": 0, "voltage_override": 0},
        {"preset_name": "275kV Monopole", "x_offset": 20, "load": 100, "arrangement": "ABC-ABC",
         "current_override": 900, "voltage_override": 0},
    ],
}


def test_default_config_solves_and_serialises():
    cfg = service.default_config()
    sol = service.solve(cfg)
    json.dumps(sol)
    assert sol["peak_b"] > 0 and sol["peak_e"] > 0 and sol["overall"] in ("PASS", "MARGINAL", "FAIL")
    assert len(sol["profile"]["x"]) == len(sol["profile"]["b0"]) == len(sol["profile"]["eS"])
    assert sol["governing"]["unit"] in ("µT", "kV/m")


def test_normalise_survives_rubbish():
    for bad in (None, 3, "x", [], {"lines": "no"}, {"lines": [1, 2]}, {"shield": 7, "buildings": {}}):
        cfg = service.normalise_config(bad)
        assert service.solve(cfg)["peak_b"] > 0


def test_legacy_scenario_is_migrated():
    cfg = service.normalise_config(copy.deepcopy(LEGACY))
    assert cfg["schema"] == service.SCHEMA and len(cfg["lines"]) == 2
    assert cfg["corridor"]["max_sag_m"] == 2.0 and cfg["corridor"]["row_half_width"] == 15.0
    assert cfg["lines"][0]["load_pct"] == 60 and cfg["lines"][1]["current_a"] == 900
    sol = service.solve(cfg)
    assert [ln["operating_a"] for ln in sol["lines"]][1] == pytest.approx(900.0)
    assert sol["lines"][0]["operating_a"] == pytest.approx(0.6 * sol["lines"][0]["rated_a"])


def test_old_standard_ids_still_load():
    cfg = service.default_config()
    cfg["standards"] = ["MS_2332_1_2009", "SI_IT_RESIDENTIAL", "nonsense", "ICNIRP_2010", "ICNIRP_2010"]
    assert service.normalise_config(cfg)["standards"] == ["MY_ICNIRP_1998", "SI_ZONE_I", "ICNIRP_2010"]


def test_missing_preset_falls_back_with_a_warning():
    cfg = service.default_config()
    cfg["lines"][0]["preset"] = "A tower that was removed"
    sol = service.solve(service.normalise_config(cfg))
    assert any("not available" in w for w in sol["warnings"])


def test_points_are_exact_and_follow_the_span():
    cfg = service.default_config()
    r = service.points(cfg, [{"x": 10, "y": 1, "z": 0, "label": "mid"}, {"x": 10, "y": 1, "z": 140, "label": "tower"}])
    mid, tower = r["rows"]
    assert tower["b0"] < mid["b0"]          # conductors are higher near the tower
    assert mid["b0"] == pytest.approx(mid["bS"])


def test_what_changed_shows_the_loading_fix():
    cfg = service.default_config()
    cfg["lines"][0]["load_pct"] = 50.0
    items = {i["id"]: i for i in service.what_changed(service.normalise_config(cfg))["items"]}
    assert items["loading"]["applies"] and items["loading"]["before"] > 1.6 * items["loading"]["after"]
    assert items["span"]["after"] < items["span"]["before"]
    assert items["earth"]["before"] >= items["earth"]["after"]


def test_templates_are_valid_projects():
    ts = service.templates()
    assert len(ts) >= 5
    for t in ts:
        assert t["summary"]["peak_b"] > 0
        assert service.normalise_config(t["config"]) == t["config"]


def test_sweeps_and_assistant():
    cfg = next(t for t in service.templates() if t["id"] == "school")["config"]
    for kind in ("material", "thickness", "height", "layers", "bonding", "geometry", "coverage", "distance", "building"):
        r = service.shield_sweep(cfg, kind)
        assert r["rows"], kind
        json.dumps(r)
    a = service.shield_assistant(cfg, target_b=1.0)
    assert a["status"] in ("already", "reachable", "not_reachable")
    m = service.shield_models(cfg)
    assert [row["model"] for row in m["rows"]] == ["physical", "analytical", "empirical"]


def test_twin_payloads_decode():
    import base64
    cfg = next(t for t in service.templates() if t["id"] == "school")["config"]
    sc = twin_service.scene(cfg)
    g = np.frombuffer(base64.b64decode(sc["ground"]["B0"]), "<f4").reshape(sc["ground"]["nz"], sc["ground"]["nx"])
    mid, end = g[sc["ground"]["nz"] // 2], g[0]
    assert mid.max() > end.max() > 0          # the ground map is strongest at mid-span
    assert mid.max() == pytest.approx(sc["peak_b"], rel=0.02)
    sec = twin_service.section(cfg, 0.0)
    assert len(base64.b64decode(sec["BS"])) == 4 * sec["nx"] * sec["ny"]
    vol = twin_service.volume(cfg)
    assert len(base64.b64decode(vol["B0"])) == vol["nx"] * vol["ny"] * vol["nz"]
    json.dumps(sc)
    # buildings carry the average over their inside (what the labels show) beside the probe-point value
    b, inside = sc["buildings"][0], service.site_for(cfg).building_stats(0)
    assert b["b_in0"] == pytest.approx(inside["b"]["avg0"], abs=1e-6) and b["b_inS"] == pytest.approx(inside["b"]["avgS"], abs=1e-6)
    assert b["e_in0"] == pytest.approx(inside["e"]["avg0"], abs=1e-6) and b["b_in0"] != b["b0"]


def test_field_lines_and_earth_envelope():
    cfg = service.default_config()
    fl = service.fieldlines(cfg)
    a = np.frombuffer(base64.b64decode(fl["a0_re"]), "<f4")
    assert a.size == fl["nx"] * fl["ny"] and fl["scale"] > 0 and fl["has_current"]
    assert fl["aS_re"] is None and not fl["shield_included"]          # no shield in the default project
    env = service.earth_envelope(cfg)
    assert max(env["perfect"]) >= max(env["free"]) > 0


def test_reports_render_in_every_format():
    cfg = next(t for t in service.templates() if t["id"] == "school")["config"]
    user = {"name": "Tester", "organisation": "Lab", "is_guest": False}
    pdf, media, name = report_service.render(cfg, "pdf", {"project": {"name": "Unit test"}}, user)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 50_000 and name.endswith(".pdf")
    docx, _, _ = report_service.render(cfg, "docx", {}, user)
    assert docx[:2] == b"PK" and len(docx) > 20_000
    txt, _, _ = report_service.render(cfg, "txt", {}, user)
    assert b"EMF ASSESSMENT REPORT" in txt and "µT".encode() in txt
    prev = report_service.preview(cfg, {}, user)
    json.dumps(prev)                                   # no raw image bytes in the preview
    assert prev["doc"]["blocks"] and "EXECUTIVE SUMMARY" in prev["text"]


def test_validation_catalogue_and_comparison():
    cat = validation_service.catalogue()
    assert cat["sources"] and cat["bibliography"]
    ds = cat["sources"][0]["datasets"][0]
    res = validation_service.compare(service.default_config(), dataset_id=ds["id"], use_published_geometry=True)
    json.dumps(res)
    csv_text = "distance_m,b_uT\n-20,1.0\n0,2.0\n20,1.1\n"
    res = validation_service.compare(service.default_config(), csv_text=csv_text)
    assert res["n_points"] == 3


def test_report_optional_sections_and_new_fields():
    cfg = next(t for t in service.templates() if t["id"] == "school")["config"]
    user = {"name": "Tester", "organisation": "Lab", "is_guest": False}
    opts = {"sections": {"options": True, "scenarios": True}, "objective": "Can a school stand 30 m from the line?",
            "scenarios": [{"name": "Base case", "config": service.default_config()}, "rubbish", {"name": "no config"}]}
    doc = report_service.build_doc(cfg, dict(opts, with_figures=False), user)
    heads = [b["text"] for b in doc["blocks"] if b["type"] == "heading" and b.get("level", 1) == 1]
    for want in ("Executive summary", "Objective", "Field results", "Standards assessment",
                 "Shielding options compared", "Scenario comparison", "Conclusion", "References"):
        assert want in heads, want
    assert heads.index("Objective") == heads.index("Executive summary") + 1
    assert heads.index("Conclusion") < heads.index("References")
    txt = report_service.preview(cfg, opts, user)["text"]
    assert "Can a school stand 30 m" in txt and "Base case" in txt and "Raise the conductors by 5 m" in txt
    assert "Highest outside the right-of-way" in txt and "not a compliance credit" in txt
    # defaults: no objective text and no optional sections -> those headings are absent
    plain = [b["text"] for b in report_service.build_doc(cfg, {"with_figures": False}, user)["blocks"] if b["type"] == "heading"]
    assert "Objective" not in plain and "Shielding options compared" not in plain and "Scenario comparison" not in plain
    assert "Field results" in plain and "Conclusion" in plain
    cat = report_service.options_catalogue()
    assert [s["id"] for s in cat["sections"]][:2] == ["summary", "objective"] and all("hint" in s for s in cat["sections"])


def test_survey_can_be_compared_with_the_shielded_field():
    cfg = next(t for t in service.templates() if t["id"] == "barrier")["config"]
    csv_text = "distance_m,b_uT\n25,1.0\n30,0.8\n40,0.5\n"
    plain = validation_service.compare(cfg, csv_text=csv_text)
    shielded = validation_service.compare(cfg, csv_text=csv_text, with_shield=True)
    assert "with its shield" in shielded["basis"] and "with its shield" not in plain["basis"]
    assert [p["taki"] for p in shielded["points"]] != [p["taki"] for p in plain["points"]]
    off = copy.deepcopy(cfg); off["shield"]["enabled"] = False
    r = validation_service.compare(service.normalise_config(off), csv_text=csv_text, with_shield=True)
    assert any("No shield is switched on" in n for n in r["notes"])
    assert [p["taki"] for p in r["points"]] == [p["taki"] for p in plain["points"]]


def test_storey_height_is_kept_for_the_drawing_only():
    cfg = service.default_config()
    cfg["buildings"][0]["floor_h"] = 1.0                       # clamped
    a = service.normalise_config(cfg)
    assert a["buildings"][0]["floor_h"] == 2.4
    cfg["buildings"][0]["floor_h"] = 5.0
    b = service.normalise_config(cfg)
    assert twin_service.scene(b)["buildings"][0]["floor_h"] == 5.0
    cfg["buildings"][0]["floor_h"] = None
    c = service.normalise_config(cfg)
    assert twin_service.scene(c)["buildings"][0]["floor_h"] > 0
    assert service.solve(b)["peak_b"] == service.solve(c)["peak_b"]              # no effect on the physics


# --------------------------------------------------------------------------- v4.2
def _school():
    return copy.deepcopy(next(t for t in service.templates() if t["id"] == "school")["config"])


def _grid(fl, key):
    return np.frombuffer(base64.b64decode(fl[key]), "<f4").reshape(fl["ny"], fl["nx"]).astype(float)


def test_field_lines_carry_the_potential_without_and_with_the_shield():
    cfg = _school()
    fl = service.fieldlines(cfg)
    json.dumps(fl)
    assert fl["shield_included"] and fl["has_current"] and fl["scale"] > 0
    from engine import field_lines as flm
    assert flm.SOFTNESS_MIN <= fl["softness"] <= flm.SOFTNESS
    a0 = _grid(fl, "a0_re") + 1j * _grid(fl, "a0_im")
    a1 = _grid(fl, "aS_re") + 1j * _grid(fl, "aS_im")
    assert np.all(np.isfinite(a0)) and np.all(np.isfinite(a1))
    assert float(np.max(np.abs(a1 - a0))) > 0.01 * fl["scale"]          # the shield's own field is in it
    # the grid is the one the engine defines, in µWb/m, over the reported window
    from engine import field_lines
    site = service.site_for(cfg)
    x, y, want = field_lines.potential_grid(site.conductors_at(fl["z"]), (fl["x0"], fl["x1"]), (fl["y0"], fl["y1"]),
                                            site.ground_model, site.earth_rho, site.freq, nx=fl["nx"], ny=fl["ny"])
    assert (x[0], x[-1], y[-1]) == pytest.approx((fl["x0"], fl["x1"], fl["y1"]))
    assert np.allclose(a0, want * 1e6, rtol=1e-4, atol=1e-6 * fl["scale"])
    off = _school(); off["shield"]["enabled"] = False
    plain = service.fieldlines(service.normalise_config(off))
    assert plain["aS_re"] is None and plain["aS_im"] is None and not plain["shield_included"]
    idle = _school()
    for ln in idle["lines"]:
        ln["load_pct"] = 0.0
    assert not service.fieldlines(service.normalise_config(idle))["has_current"]


def test_field_line_helpers():
    from engine import field_lines as fl
    # one wire carrying current out of the page: A = -k ln r, and the field circles it anticlockwise
    x = np.linspace(-10.0, 10.0, 81); y = np.linspace(-10.0, 10.0, 81)
    X, Y = np.meshgrid(x, y)
    a = -np.log(np.hypot(X, Y) + 1e-9)
    bx, by = fl.direction(a, x, y)
    j, i = 40, 60                                            # a point on the +x axis
    assert abs(bx[j, i]) < 1e-9 and by[j, i] > 0
    j, i = 60, 40                                            # a point on the +y axis
    assert bx[j, i] < 0 and abs(by[j, i]) < 1e-9
    arr = fl.arrows(a, x, y)
    assert np.allclose(np.hypot(arr["u"], arr["v"]), 1.0) and len(arr["x"]) == len(arr["u"]) > 20
    assert fl.arrows(np.zeros_like(a), x, y) is None
    # the spreading used for "show weak field": odd, monotonic, within -1..1, linear near zero
    v = np.linspace(-5.0, 5.0, 201)
    s = fl.spread(v, 1.0)
    assert np.all(np.diff(s) > 0) and np.allclose(s, -s[::-1]) and s[100] == 0.0
    assert fl.spread(np.array([1.0]), 1.0)[0] == pytest.approx(1.0) and np.all(fl.spread(v, 0.0) == 0.0)
    small = fl.spread(np.array([1e-6, 2e-6]), 1.0)
    assert small[1] / small[0] == pytest.approx(2.0, rel=1e-3)
    # a snapshot a quarter of a cycle later is the other part of the phasor
    ph = np.array([1.0 + 2.0j])
    assert fl.snapshot(ph, 0.0)[0] == pytest.approx(1.0) and fl.snapshot(ph, 90.0)[0] == pytest.approx(-2.0)
    assert fl.scale_of(np.array([np.nan, 1.0, 2.0, 3.0])) == pytest.approx(np.percentile([1.0, 2.0, 3.0], 99.0))
    assert fl.scale_of(np.array([])) == 0.0
    # how far down the spacing reaches follows the picture: the weakest fifth sets the linear part
    far = np.linspace(0.0, 1.0, 1001) ** 2 * np.exp(0.7j)     # the weakest fifth lies below 0.2^2 = 0.04
    assert fl.softness_of(far) == pytest.approx(fl.scale_of(far) / 0.04, rel=1e-6)
    assert fl.softness_of(far ** 4) == fl.SOFTNESS            # a field that dies away quickly: as soft as allowed
    assert fl.softness_of(np.ones(50)) == fl.SOFTNESS_MIN and fl.softness_of(np.zeros(50)) == fl.SOFTNESS
    assert fl.softness_of(far * np.exp(1.3j)) == pytest.approx(fl.softness_of(far))   # the instant does not matter
    wide, narrow = fl.spread(np.array([0.01]), 1.0, 400.0)[0], fl.spread(np.array([0.01]), 1.0, 20.0)[0]
    assert wide > 2.5 * narrow                                # a softer mapping lifts weak values further


def _tables(doc):
    out, title = {}, ""
    for b in doc["blocks"]:
        if b["type"] == "heading":
            title = b["text"]
        elif b["type"] == "table":
            out.setdefault(title, b)
    return out


USER = {"name": "Tester", "organisation": "Lab", "is_guest": False}


def test_report_without_a_shield_has_no_with_shield_columns():
    cfg = _school()
    cfg["points"] = [{"x": 20.0, "y": 1.0, "z": 0.0, "label": "Gate"}, {"x": 40.0, "y": 1.0, "z": 400.0, "label": "Far"}]
    cfg = service.normalise_config(cfg)
    on = _tables(report_service.build_doc(cfg, {"with_figures": False}, USER))
    pts = next(t for h, t in on.items() if "point" in h.lower())
    assert "B with shield" in pts["header"] and "E with shield" in pts["header"]
    assert len(pts["header"]) == len(pts["widths"]) == len(pts["rows"][0]) == 9
    assert pts["rows"][1][6] == "beyond the shield"                         # outside the shield's length
    assert "%" in pts["rows"][0][6]
    off = copy.deepcopy(cfg); off["shield"]["enabled"] = False
    doc = report_service.build_doc(service.normalise_config(off), {"with_figures": False}, USER)
    plain = next(t for h, t in _tables(doc).items() if "point" in h.lower())
    assert plain["header"] == ["#", "Label", "x (m)", "y (m)", "z (m)", "B (µT)", "E (kV/m)"]
    assert all(len(r) == 7 for r in plain["rows"]) and len(plain["widths"]) == 7
    text = report_service.preview(service.normalise_config(off), {}, USER)["text"]
    assert "with shield" not in text.lower() and "beyond the shield" not in text


def test_report_references_are_listed_once():
    cfg = service.default_config()
    cfg["standards"] = ["ICNIRP_2010", "ICNIRP_1998", "MY_ICNIRP_1998", "ICNIRP_1998_OCC", "EU_1999_519_EC"]
    cfg = service.normalise_config(cfg)
    assert len(cfg["standards"]) == 5
    doc = report_service.build_doc(cfg, {"with_figures": False}, USER)
    i = next(k for k, b in enumerate(doc["blocks"]) if b["type"] == "heading" and b["text"] == "References")
    refs = doc["blocks"][i + 1]["items"]
    assert len(refs) >= 6 and len(refs) == len(set(refs))
    for mark in ("99(6)", "74(4)"):                           # the two ICNIRP guidelines, by volume(issue)
        assert sum(mark in r for r in refs) == 1, mark
    y98 = next(r for r in refs if "74(4)" in r)               # one entry, under every standard that rests on it
    names = y98.split(": ICNIRP.")[0]
    assert names == "ICNIRP 1998; ICNIRP 1998 (occupational); Malaysia (ICNIRP 1998 basis)"
    assert y98.endswith("Applied in Malaysian practice (Ministry of Health / Suruhanjaya Tenaga).")
    assert sum("Deri" in r for r in refs) == 1 and sum("IEEE Std 644" in r for r in refs) == 1


def test_report_describes_circuits_sag_and_clearance():
    cfg = next(t for t in service.templates() if t["id"] == "quad")["config"]
    doc = report_service.build_doc(cfg, {"with_figures": False}, USER)
    tabs = _tables(doc)
    conf = dict((r[0], r[1]) for r in next(iter(tabs.values()))["rows"])
    line = conf[cfg["lines"][0]["name"]]
    assert "upper left circuit: 275 kV" in line and "out of service (earthed)" in line and "lower" in line
    assert "Sag and span" in conf and "Lowest conductor at mid-span" in conf and "Thermal sag at 100% load" in conf
    assert "independently of the span" in conf["Sag and span"]
    cond = next(t for h, t in tabs.items() if "conductor" in h.lower())
    off = [r for r in cond["rows"] if r[6] == "off"]
    assert len(off) == 3 and all(r[7] == "earthed" for r in off)
    text = " ".join(report_service.preview(cfg, {}, USER)["text"].split())
    assert "de-energised and earthed" in text and "not modelled" in text
    follow = copy.deepcopy(cfg); follow["corridor"].update({"sag_follows_span": True, "span_m": 360.0})
    conf2 = dict((r[0], r[1]) for r in next(iter(_tables(report_service.build_doc(
        service.normalise_config(follow), {"with_figures": False}, USER)).values()))["rows"])
    assert "sag follows the span" in conf2["Sag and span"] and "1.44" in conf2["Sag and span"]
    assert "300 m reference span" in conf2["Thermal sag at 100% load"]


def test_report_judges_the_shield_inside_the_building():
    cfg = _school()
    doc = report_service.build_doc(cfg, {"with_figures": False, "sections": {"options": True}}, USER)
    heads = [b["text"] for b in doc["blocks"] if b["type"] == "heading"]
    name = cfg["buildings"][cfg["shield"]["target_building"]]["name"]
    assert f"The same shield under each model (average inside {name})" in heads
    assert f"Arrangements and measures at the line (average inside {name})" in heads
    tabs = _tables(doc)
    shield = dict((r[0], r[1]) for r in next(t for h, t in tabs.items() if h.lower().startswith("shield"))["rows"])
    assert any(k.startswith("Magnetic field inside") for k in shield)
    assert "At one point: 1 m inside the wall facing the line" in shield
    opts = next(t for h, t in tabs.items() if h.startswith("Arrangements"))
    assert opts["header"] == ["Option", "B inside", "B change", "B worst point", "E inside", "E change"]
    assert any("m from the centreline" in r[0] for r in opts["rows"])       # barrier rows say where they stand
    rec = next(t for h, t in tabs.items() if "building" in h.lower() or "receptor" in h.lower())
    assert rec["header"][2:5] == ["B inside, average", "with the shield", "E inside, average"]
    assert len(rec["header"]) == len(rec["widths"]) == len(rec["rows"][0])
    ctx = report_service.ai_context(cfg)
    assert any("average" in s.lower() and "inside" in s.lower() for s in ctx["shield"])


def test_report_small_numbers_and_units():
    assert [report_service._pct_small(v) for v in (0.0, 1e-9, -3e-4, 0.004, 0.25, 3.14159)] == \
           ["0%", "0%", "0%", "0.004%", "0.25%", "3.1%"]
    assert report_service._lo(1e-12) == "about 0" and report_service._lo(0.002, 5.0) == "about 0"
    assert report_service._lo(0.1234) == "0.123" and report_service._lo(None) == "-"
    from engine import report
    nb = " "
    assert report._nb("2.5 µT and 0.4 kV/m at 30 m") == f"2.5{nb}µT and 0.4{nb}kV/m at 30{nb}m"
    assert report._nb("132 kV, 600 A, 50 Hz, 3 mm") == f"132{nb}kV, 600{nb}A, 50{nb}Hz, 3{nb}mm"
    assert report._nb("x = 12.5 m") == f"x{nb}={nb}12.5{nb}m"
    assert report._nb("No numbers here") == "No numbers here" and report._nb("") == ""
    assert report._nb("in 2010 the limit") == "in 2010 the limit"         # a year is not a quantity
    doc = report_service.build_doc(service.default_config(), {"with_figures": False,
                                                              "sections": {"validation": True}}, USER)
    val = next(t for t in _tables(doc).values() if t["header"][0] == "Check")
    assert len(val["rows"]) >= 11 and all(r[4] == "PASS" for r in val["rows"])
    assert all("ratio" not in c for r in val["rows"] for c in r)
    assert all(r[3].endswith("%") and "e-" not in r[3] for r in val["rows"])


def test_validation_says_when_a_comparison_is_not_like_for_like():
    cfg = service.default_config()
    ds = "275_132kV_QC__untransposed__cond"                   # a tower whose geometry the source published
    own = validation_service.compare(cfg, dataset_id=ds)
    same = validation_service.compare(cfg, dataset_id=ds, use_published_geometry=True)
    assert own["like_for_like"] is False and "current project" in own["basis"]
    assert same["like_for_like"] is True and "published tower geometry" in same["basis"]
    assert same["mean_ratio"] == pytest.approx(2 ** 0.5, rel=0.02)          # the known peak / RMS convention
    # the source gives no geometry for its 132 kV tower: asking for it cannot make the comparison fair
    none = validation_service.compare(cfg, dataset_id="132kV_DC__untransposed__cond", use_published_geometry=True)
    assert none["like_for_like"] is False and any("did not publish tower geometry" in n for n in none["notes"])
    survey = validation_service.compare(cfg, csv_text="distance_m,b_uT\n-20,1.0\n0,2.0\n20,1.1\n")
    assert survey["like_for_like"] is None


def test_scenarios_report_what_each_shield_does():
    on = _school()
    off = copy.deepcopy(on); off["shield"]["enabled"] = False
    rows = service.compare_scenarios([{"name": "with", "config": on},
                                      {"name": "without", "config": service.normalise_config(off)}])["rows"]
    eff = rows[0]["shield_effect"]
    head = service.shield_headline(service.site_for(on))
    assert eff["basis"] == "inside" and eff["b_red_pct"] == pytest.approx(head["b_red_pct"])
    assert eff["e_red_pct"] == pytest.approx(head["e_red_pct"])
    assert rows[1]["shield"] == "off" and not rows[1].get("shield_effect")
    assert rows[0]["peak_b"] == rows[1]["peak_b"]             # peaks are unshielded values either way


def test_what_changed_follows_each_circuits_loading():
    cfg = service.default_config()
    cfg["lines"][0]["circuits"] = [{"on": True, "load_pct": 40.0}, {"on": True, "load_pct": 40.0}]
    items = {i["id"]: i for i in service.what_changed(service.normalise_config(cfg))["items"]}
    assert items["loading"]["applies"] and items["loading"]["before"] > 2.0 * items["loading"]["after"]
    full = service.default_config()
    full["lines"][0]["load_pct"] = 100.0
    same = {i["id"]: i for i in service.what_changed(service.normalise_config(full))["items"]}
    assert same["loading"]["before"] == pytest.approx(same["loading"]["after"], rel=0.02)


def test_word_report_tables_follow_the_same_layout_as_the_pdf():
    import io
    from docx import Document
    data, _, _ = report_service.render(_school(), "docx", {"with_figures": False, "sections": {"options": True}}, USER)
    doc = Document(io.BytesIO(data))
    opts = next(t for t in doc.tables if t.rows[0].cells[0].text == "Option")
    widths = [c.width for c in opts.rows[1].cells]
    assert widths[0] > 2.5 * widths[1] and abs(sum(widths) / 360000 - 17.0) < 0.2     # the wide first column; 17 cm in all
    assert "w:tblHeader" in opts.rows[0]._tr.xml and "w:cantSplit" in opts.rows[1]._tr.xml
    small = next(t for t in doc.tables if t.rows[0].cells[0].text == "Model")
    assert all(p.paragraph_format.keep_with_next for r in small.rows[:-1] for c in r.cells for p in c.paragraphs)
    assert not small.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next


# --------------------------------------------------------------------------- v4.2.1
def test_electric_field_lines_carry_charges_potential_and_the_shield():
    cfg = _school()
    el = service.efieldlines(cfg)
    json.dumps(el)
    site = service.site_for(cfg)
    conds = site.conductors_at(el["z"])
    assert el["shield_included"] and el["has_voltage"] and el["lam_ref"] > 0 and el["scale"] > 0
    assert len(el["conductors"]) == len(el["lam0"]) == len(el["lamS"]) == len(conds)
    assert el["walls"] == [[round(v, 4) for v in w] for w in site.geom.walls]
    # the charges are the ones the electric field everywhere else is computed from, as q / (2 pi eps0) in kV
    from engine import field_lines, physics
    q = physics.solve_conductor_charges(conds, bundle_equivalent=site.bundle_eq)
    lam0 = np.array([complex(a, b) for a, b in el["lam0"]])
    assert np.allclose(lam0, q * field_lines.K_E / 1e3, rtol=1e-5, atol=1e-6)
    assert el["lam_ref"] == pytest.approx(float(np.max(np.abs(lam0))), rel=1e-5)
    # the potential is zero on the ground and reaches the phase voltage at the wires
    v0 = _grid(el, "v0_re") + 1j * _grid(el, "v0_im")
    vS = _grid(el, "vS_re") + 1j * _grid(el, "vS_im")
    assert np.all(np.isfinite(v0)) and np.all(np.isfinite(vS))
    assert float(np.max(np.abs(v0[0]))) < 1e-6 and float(np.max(np.abs(vS[0]))) < 1e-3
    assert float(np.max(np.abs(vS - v0))) > 0.01 * el["scale"]             # the charge on the shield is in it
    off = _school(); off["shield"]["enabled"] = False
    plain = service.efieldlines(service.normalise_config(off))
    assert plain["lamS"] is None and plain["vS_re"] is None and plain["exS_re"] is None and not plain["shield_included"]
    assert plain["walls"] == [] and plain["wires"] == []


def test_electric_potential_and_field_agree_and_the_shield_is_an_equipotential():
    """E = -grad V for what the browser is sent, and an earthed sheet sits at the voltage of the ground."""
    from engine import field_lines, physics, shield_bem
    cfg = _school()
    site = service.site_for(cfg)
    conds = site.conductors_at(0.0)
    _, ele = site.solution(0.0)
    d = field_lines.electric_data(conds, (site.x_min, site.x_max), (0.0, site.y_max), site.bundle_eq, ele)
    x, y = d["x"], d["y"]
    X, Y = np.meshgrid(x, y)
    # without the shield: the summed line charges give the field every other page reports
    ex, ey = field_lines.electric_field(conds, d["lam0"], X, Y)
    rms = np.hypot(np.abs(ex), np.abs(ey)) / 1e3
    assert np.allclose(rms, physics.compute_e_field_at_points(conds, X, Y, bundle_equivalent=site.bundle_eq), rtol=1e-9, atol=1e-12)
    # with it: wires plus the induced part add up to the solver's field
    exc, eyc = field_lines.electric_field(conds, d["lamS"], X, Y)
    tot = np.hypot(np.abs(exc + d["exS"]), np.abs(eyc + d["eyS"])) / 1e3
    assert np.allclose(tot, ele.e_rms_kVm(X, Y), rtol=1e-9, atol=1e-12)
    # the gradient of the potential is the field, away from the wires and the sheets where a grid cannot follow it
    clear = np.ones(X.shape, bool)
    for c in conds:
        clear &= np.hypot(X - c.x, Y - c.y_sagged) > 5.0
    for cx, cy in zip(site._mesh.cx, site._mesh.cy):
        clear &= np.hypot(X - cx, Y - cy) > 3.0
    clear[:2] = clear[-2:] = False; clear[:, :2] = clear[:, -2:] = False
    for v, fx, fy in ((d["v0"], ex, ey), (d["vS"], exc + d["exS"], eyc + d["eyS"])):
        gy, gx = np.gradient(v, y, x)
        err = np.hypot(np.abs(fx + gx), np.abs(fy + gy)) / np.hypot(np.abs(fx), np.abs(fy))
        assert float(np.median(err[clear])) < 0.005 and float(np.percentile(err[clear], 95)) < 0.03
    # every earthed element of the shield is at zero potential, against a phase voltage of some 160 kV
    m = site._mesh
    assert np.all(m.grounded)
    assert float(np.max(np.abs(ele.v_total(m.cx.copy(), m.cy.copy())))) < 1.0
    # and at the surface of each wire the potential is its phase voltage
    for c in conds:
        r = physics.bundle_equivalent_radius(c)
        v = complex(ele.v_total(np.array([c.x + r]), np.array([c.y_sagged]))[0])
        assert abs(v) == pytest.approx(c.voltage_kV * 1e3 / np.sqrt(3.0), rel=5e-3)
    # an unearthed barrier floats at one voltage of its own
    fl = _school(); fl["shield"].update({"preset": "wall-between", "grounded": False})
    s2 = service.site_for(service.normalise_config(fl))
    _, e2 = s2.solution(0.0)
    vp = e2.v_total(s2._mesh.cx.copy(), s2._mesh.cy.copy())
    (vf,) = e2.v_float.values()
    assert abs(vf) > 10.0 and np.allclose(vp, vf, rtol=0, atol=1e-3 * abs(vf))

