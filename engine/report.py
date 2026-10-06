"""
report.py
=========
Renders a report "document model" to PDF (reportlab), Word (python-docx) and
plain text. The model is a list of blocks, so the three formats always carry
the same content:

    {"type": "heading", "text": ..., "level": 1|2}
    {"type": "para", "text": ...}
    {"type": "note", "text": ...}                      small grey caption
    {"type": "status", "text": ..., "status": "PASS"|"MARGINAL"|"FAIL"|"NOT_ASSESSED"}
    {"type": "table", "header": [...], "rows": [[...]], "widths": [...]?, "status_col": int?}
    {"type": "figure", "title": ..., "png": bytes | None, "caption": ...?}
    {"type": "bullets", "items": [...]}
    {"type": "pagebreak"}

The model itself is assembled in server/report_service.py from the solved site.
AI-drafted text, when present, arrives as ordinary blocks under a heading that
labels it as AI-generated; nothing in this module decides compliance.
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Dict, List, Optional

STATUS_HEX = {"PASS": "#1E6B3F", "MARGINAL": "#B26B00", "FAIL": "#B3261E", "NOT_ASSESSED": "#5B6770"}
STATUS_LABEL = {"PASS": "Compliant", "MARGINAL": "Compliant - low margin", "FAIL": "Exceeds limit",
                "NOT_ASSESSED": "No applicable limit"}

#: In the PDF a table with up to this many rows stays on one page with its heading.
KEEP_ROWS = 14

#: Width of the text area of the Word page: A4 less the 2 cm margins set in build_docx.
PAGE_TEXT_CM = 17.0


#: Characters the PDF's built-in Helvetica cannot draw, and what to print instead.
_PDF_MAP = str.maketrans({
    "\u2212": "-", "\u2248": "~", "\u2192": "->", "\u2264": "<=", "\u2265": ">=", "\u03bc": "\u00b5",
    "\u03a9": "ohm", "\u03c1": "rho", "\u03c3": "sigma", "\u221a": "sqrt", "\u221e": "inf",
    "\u2082": "2", "\u2080": "0", "\u00b7": "\u00b7", "\u2009": " ", "\u202f": " ", "\u03c9": "w",
    "\u0394": "delta ", "\u03c0": "pi", "\u2260": "!=", "\u2713": "ok", "\u2717": "x",
})


#: A number and its unit stay on one line ("2.129 kV/m", "x = 26.0 m", "(-18%)").
_UNIT_GAP = re.compile(r"(?<=\d) (?=(?:µT|\u00b5T|kV/m|kV|mm2|mm|m|A|dB|Hz|W|S/m|V)(?![A-Za-z]))")
_EQ_GAP = re.compile(r"(?<=[xyz]) = (?=[-+\u2212]?\d)")


def _nb(s) -> str:
    """Join a number to its unit with a no-break space so table cells do not split them."""
    s = _UNIT_GAP.sub("\u00a0", str(s))
    return _EQ_GAP.sub("\u00a0=\u00a0", s)


def _esc(s) -> str:
    s = _nb(s).translate(_PDF_MAP)
    s = s.encode("cp1252", "replace").decode("cp1252")
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def build_pdf(doc: Dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (CondPageBreak, Image, KeepTogether, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)

    PRIMARY = colors.HexColor("#004B87")
    INK = colors.HexColor("#0F1B24")
    STEEL = colors.HexColor("#5B6770")
    LINE = colors.HexColor("#DDE2E6")
    SURF = colors.HexColor("#F7F9FA")

    buf = io.BytesIO()
    page_w, page_h = A4
    margin = 1.9 * cm
    frame_w = page_w - 2 * margin
    title = doc.get("title", "EMF assessment report")
    project = doc.get("project", "")

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setStrokeColor(LINE)
        canvas.setLineWidth(0.6)
        canvas.line(margin, page_h - 1.25 * cm, page_w - margin, page_h - 1.25 * cm)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.setFillColor(PRIMARY)
        canvas.drawString(margin, page_h - 1.05 * cm, "Taki")
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(STEEL)
        canvas.drawString(margin + 1.0 * cm, page_h - 1.05 * cm, "EMF simulation report")
        canvas.drawRightString(page_w - margin, page_h - 1.05 * cm, project[:70])
        canvas.line(margin, 1.25 * cm, page_w - margin, 1.25 * cm)
        canvas.drawString(margin, 0.85 * cm, doc.get("date", ""))
        canvas.drawRightString(page_w - margin, 0.85 * cm, f"Page {d.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=margin, rightMargin=margin,
                            topMargin=1.9 * cm, bottomMargin=1.8 * cm, title=title,
                            author=doc.get("author", "") or "Taki")
    ss = getSampleStyleSheet()
    st_title = ParagraphStyle("T", parent=ss["Title"], textColor=PRIMARY, fontSize=22, leading=26,
                              alignment=0, spaceAfter=4)
    st_sub = ParagraphStyle("S", parent=ss["BodyText"], textColor=INK, fontSize=11, leading=14)
    st_meta = ParagraphStyle("M", parent=ss["BodyText"], textColor=STEEL, fontSize=8.5, leading=11)
    st_h1 = ParagraphStyle("H1", parent=ss["Heading2"], textColor=PRIMARY, fontSize=13, leading=16,
                           spaceBefore=14, spaceAfter=6)
    st_h2 = ParagraphStyle("H2", parent=ss["Heading3"], textColor=INK, fontSize=10.5, leading=13,
                           spaceBefore=8, spaceAfter=4)
    st_body = ParagraphStyle("B", parent=ss["BodyText"], textColor=INK, fontSize=9.5, leading=13.5,
                             spaceAfter=5)
    st_note = ParagraphStyle("N", parent=ss["BodyText"], textColor=STEEL, fontSize=8, leading=10.5,
                             spaceAfter=4)
    st_cell = ParagraphStyle("C", parent=ss["BodyText"], textColor=INK, fontSize=8, leading=10)
    st_head = ParagraphStyle("CH", parent=st_cell, textColor=colors.white, fontName="Helvetica-Bold")

    story = [Paragraph(_esc(title), st_title)]
    if doc.get("subtitle"):
        story.append(Paragraph(_esc(doc["subtitle"]), st_sub))
    meta = " &nbsp;|&nbsp; ".join(_esc(x) for x in [
        doc.get("project") and f"Project: {doc['project']}",
        doc.get("author") and f"Prepared by: {doc['author']}",
        doc.get("organisation"), doc.get("date") and f"Generated: {doc['date']}"] if x)
    if meta:
        story.append(Paragraph(meta, st_meta))
    story.append(Spacer(1, 0.3 * cm))

    # Headings wait here until what they introduce arrives, and then share a page with it:
    # a heading is never left alone at the foot of a page, and a table of up to KEEP_ROWS
    # rows is never cut in two (a longer one may split, repeating its header row).
    pending: list = []

    def emit(items: list, together: bool = True) -> None:
        nonlocal pending
        if together:
            story.append(KeepTogether(pending + items) if pending or len(items) > 1 else items[0])
        else:
            if pending:
                story.append(CondPageBreak((4.2 + 0.9 * (len(pending) - 1)) * cm))
                story.extend(pending)
            story.extend(items)
        pending = []

    blocks = doc.get("blocks", [])
    taken = set()                      # notes already placed under the table they explain
    for bi, blk in enumerate(blocks):
        kind = blk.get("type")
        if bi in taken:
            continue
        if kind == "heading":
            pending.append(Paragraph(_esc(blk["text"]), st_h1 if blk.get("level", 1) == 1 else st_h2))
        elif kind == "para":
            paras = [Paragraph(_esc(para.strip()), st_body) for para in str(blk["text"]).split("\n") if para.strip()]
            if paras:
                emit(paras[:1])
                story.extend(paras[1:])
        elif kind == "note":
            emit([Paragraph(_esc(blk["text"]), st_note)])
        elif kind == "bullets":
            items = [Paragraph("&bull;&nbsp; " + _esc(it), st_body) for it in blk.get("items", [])]
            if items:
                emit(items[:1])
                story.extend(items[1:])
        elif kind == "status":
            col = colors.HexColor(STATUS_HEX.get(blk.get("status"), "#5B6770"))
            t = Table([[Paragraph(f"<b>{_esc(blk['text'])}</b>",
                                  ParagraphStyle("st", parent=st_body, textColor=col, spaceAfter=0))]],
                      colWidths=[frame_w])
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), SURF),
                                   ("LINEBEFORE", (0, 0), (0, -1), 3, col),
                                   ("BOX", (0, 0), (-1, -1), 0.5, LINE),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 9),
                                   ("TOPPADDING", (0, 0), (-1, -1), 6),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            emit([t, Spacer(1, 0.2 * cm)])
        elif kind == "table":
            header, rows = blk.get("header", []), blk.get("rows", [])
            if not rows:
                continue
            ncol = max(len(header), max(len(r) for r in rows))
            widths = blk.get("widths")
            if widths and len(widths) == ncol:
                tot = float(sum(widths))
                col_w = [frame_w * w / tot for w in widths]
            else:
                col_w = [frame_w / ncol] * ncol
            data = [[Paragraph(_esc(h), st_head) for h in header]] if header else []
            for r in rows:
                data.append([Paragraph(_esc(c), st_cell) for c in list(r) + [""] * (ncol - len(r))])
            t = Table(data, colWidths=col_w, repeatRows=1 if header else 0, splitByRow=1)
            style = [("GRID", (0, 0), (-1, -1), 0.4, LINE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                     ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                     ("ROWBACKGROUNDS", (0, 1 if header else 0), (-1, -1), [colors.white, SURF])]
            if header:
                style.append(("BACKGROUND", (0, 0), (-1, 0), PRIMARY))
            sc = blk.get("status_col")
            if sc is not None:
                for i, r in enumerate(rows):
                    key = str(r[sc]).upper().replace(" ", "_") if sc < len(r) else ""
                    if key in STATUS_HEX:
                        style.append(("TEXTCOLOR", (sc, i + (1 if header else 0)),
                                      (sc, i + (1 if header else 0)), colors.HexColor(STATUS_HEX[key])))
            t.setStyle(TableStyle(style))
            items = [t, Spacer(1, 0.15 * cm)]
            small = len(rows) <= KEEP_ROWS
            if small and bi + 1 < len(blocks) and blocks[bi + 1].get("type") == "note":
                items.append(Paragraph(_esc(blocks[bi + 1]["text"]), st_note))     # its note stays under it
                taken.add(bi + 1)
            emit(items, together=small)
        elif kind == "figure":
            items = [Paragraph(_esc(blk.get("title", "")), st_h2)]
            if blk.get("png"):
                from reportlab.lib.utils import ImageReader
                iw, ih = ImageReader(io.BytesIO(blk["png"])).getSize()
                w = frame_w
                h = w * ih / iw
                max_h = 15.5 * cm
                if h > max_h:
                    h, w = max_h, max_h * iw / ih
                items.append(Image(io.BytesIO(blk["png"]), width=w, height=h))
            else:
                items.append(Paragraph("[Figure not available.]", st_note))
            if blk.get("caption"):
                items.append(Paragraph(_esc(blk["caption"]), st_note))
            emit(items)
            story.append(Spacer(1, 0.2 * cm))
        elif kind == "pagebreak":
            story.extend(pending)
            pending = []
            story.append(PageBreak())
    story.extend(pending)

    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def build_docx(doc: Dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def rgb(hexstr):
        h = hexstr.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    PRIMARY, STEEL, INK = rgb("#004B87"), rgb("#5B6770"), rgb("#0F1B24")
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    base = d.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor.lstrip("#"))
        tcPr.append(shd)

    def row_flag(row, tag):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement(tag)
        el.set(qn("w:val"), "true")
        trPr.append(el)

    h = d.add_heading(doc.get("title", "EMF assessment report"), level=0)
    for r in h.runs:
        r.font.color.rgb = PRIMARY
    if doc.get("subtitle"):
        p = d.add_paragraph(doc["subtitle"])
        p.runs[0].font.size = Pt(11)
    meta = " | ".join(x for x in [
        doc.get("project") and f"Project: {doc['project']}",
        doc.get("author") and f"Prepared by: {doc['author']}",
        doc.get("organisation"), doc.get("date") and f"Generated: {doc['date']}"] if x)
    if meta:
        p = d.add_paragraph(meta)
        p.runs[0].font.size = Pt(8.5); p.runs[0].font.color.rgb = STEEL

    for blk in doc.get("blocks", []):
        kind = blk.get("type")
        if kind == "heading":
            hh = d.add_heading(blk["text"], level=1 if blk.get("level", 1) == 1 else 2)
            for r in hh.runs:
                r.font.color.rgb = PRIMARY if blk.get("level", 1) == 1 else INK
        elif kind == "para":
            for para in str(blk["text"]).split("\n"):
                if para.strip():
                    d.add_paragraph(para.strip())
        elif kind == "note":
            p = d.add_paragraph(blk["text"])
            p.runs[0].font.size = Pt(8.5); p.runs[0].font.color.rgb = STEEL
        elif kind == "bullets":
            for it in blk.get("items", []):
                d.add_paragraph(str(it), style="List Bullet")
        elif kind == "status":
            p = d.add_paragraph()
            r = p.add_run(blk["text"])
            r.bold = True; r.font.size = Pt(11)
            r.font.color.rgb = rgb(STATUS_HEX.get(blk.get("status"), "#5B6770"))
        elif kind == "table":
            header, rows = blk.get("header", []), blk.get("rows", [])
            if not rows:
                continue
            ncol = max(len(header), max(len(r) for r in rows))
            t = d.add_table(rows=1 if header else 0, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            if header:
                for i, hd in enumerate(header):
                    c = t.rows[0].cells[i]
                    c.text = ""
                    run = c.paragraphs[0].add_run(str(hd))
                    run.bold = True; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(255, 255, 255)
                    shade(c, "#004B87")
                row_flag(t.rows[0], "w:tblHeader")           # repeat on every page the table runs to
            sc = blk.get("status_col")
            for ri, rv in enumerate(rows):
                new_row = t.add_row()
                row_flag(new_row, "w:cantSplit")             # a row stays on one page
                cells = new_row.cells
                for i in range(ncol):
                    val = _nb(rv[i]) if i < len(rv) else ""
                    cells[i].text = ""
                    run = cells[i].paragraphs[0].add_run(val)
                    run.font.size = Pt(8.5)
                    key = val.upper().replace("\u00a0", " ").replace(" ", "_")
                    if sc is not None and i == sc and key in STATUS_HEX:
                        run.bold = True; run.font.color.rgb = rgb(STATUS_HEX[key])
                    if ri % 2:
                        shade(cells[i], "#F7F9FA")
            # the same column proportions as the PDF (left alone, every column gets the same width)
            widths = blk.get("widths")
            if widths and len(widths) == ncol:
                tot = float(sum(widths))
                col_w = [Cm(PAGE_TEXT_CM * w / tot) for w in widths]
                t.autofit = False
                for i, w in enumerate(col_w):
                    t.columns[i].width = w
                for row in t.rows:
                    for i, cell in enumerate(row.cells):
                        cell.width = col_w[i]
            if len(rows) <= KEEP_ROWS:                        # a short table is not cut in two by a page break
                for row in t.rows[:-1]:
                    for cell in row.cells:
                        for par in cell.paragraphs:
                            par.paragraph_format.keep_with_next = True
            d.add_paragraph()
        elif kind == "figure":
            hh = d.add_heading(blk.get("title", ""), level=2)
            for r in hh.runs:
                r.font.color.rgb = INK
            if blk.get("png"):
                d.add_picture(io.BytesIO(blk["png"]), width=Cm(16.5))
            else:
                d.add_paragraph("[Figure not available.]")
            if blk.get("caption"):
                p = d.add_paragraph(blk["caption"])
                p.runs[0].font.size = Pt(8.5); p.runs[0].font.color.rgb = STEEL
        elif kind == "pagebreak":
            d.add_page_break()

    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Plain text
# ---------------------------------------------------------------------------
def build_txt(doc: Dict) -> str:
    W = 78
    lines = ["=" * W, doc.get("title", "EMF assessment report").upper(), "=" * W]
    for k, lab in (("subtitle", ""), ("project", "Project: "), ("author", "Prepared by: "),
                   ("organisation", ""), ("date", "Generated: ")):
        if doc.get(k):
            lines.append(f"{lab}{doc[k]}")
    for blk in doc.get("blocks", []):
        kind = blk.get("type")
        if kind == "heading":
            if blk.get("level", 1) == 1:
                lines += ["", str(blk["text"]).upper(), "-" * W]
            else:
                lines += ["", str(blk["text"])]
        elif kind in ("para", "note"):
            import textwrap
            for para in str(blk["text"]).split("\n"):
                if para.strip():
                    lines += textwrap.wrap(para.strip(), W) or [""]
        elif kind == "bullets":
            import textwrap
            for it in blk.get("items", []):
                wrapped = textwrap.wrap(str(it), W - 4) or [""]
                lines.append("  - " + wrapped[0])
                lines += ["    " + w for w in wrapped[1:]]
        elif kind == "status":
            lines.append(f">> {blk['text']}")
        elif kind == "table":
            header, rows = blk.get("header", []), blk.get("rows", [])
            if not rows:
                continue
            table = ([header] if header else []) + [list(map(str, r)) for r in rows]
            ncol = max(len(r) for r in table)
            table = [list(r) + [""] * (ncol - len(r)) for r in table]
            wid = [min(34, max(len(str(r[i])) for r in table)) for i in range(ncol)]
            for ri, r in enumerate(table):
                lines.append("  " + "  ".join(str(c)[:wid[i]].ljust(wid[i]) for i, c in enumerate(r)).rstrip())
                if header and ri == 0:
                    lines.append("  " + "  ".join("-" * w for w in wid))
        elif kind == "figure":
            lines.append(f"[Figure: {blk.get('title', '')} - see the PDF or Word report]")
    lines += ["", "=" * W]
    return "\n".join(lines)


def stamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")
